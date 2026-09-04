"""Pipeline A: xlsx_robust — main workhorse for Excel parsing.

Based on SST-Cube Extractor v3 Pipeline A.
Supports multi-level headers, CN/EN bilingual rows, unit extraction,
auto-detection of cross-section vs time-series tables.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from src.schema import Record

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


_UNIT_PATTERNS = [
    r"\((千元|万元|百万元|亿元|元|美元|万美元|亿美元|人|万人|个|件|%)\)",
    r"\uff08(千元|万元|百万元|亿元|元|美元|万美元|亿美元|人|万人|个|件|%)\uff09",
    r"单位[\uff1a:](千元|万元|百万元|亿元|元|美元|万美元|亿美元|人|万人|个|件|%)",
]

_REGION_KEYWORDS = {"北京", "天津", "上海", "重庆", "河北", "山西", "辽宁", "吉林",
    "黑龙江", "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北",
    "湖南", "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海",
    "内蒙古", "广西", "西藏", "宁夏", "新疆", "全国", "总计", "合计"}


def extract_unit(text: str) -> str:
    """Extract unit from header text like '营业收入(千元)'."""
    for pattern in _UNIT_PATTERNS:
        m = re.search(pattern, text)
        if m:
            return m.group(1)
    return ""


def split_cn_en(text: str) -> tuple[str, str]:
    """Split bilingual header like 'GDP 国内生产总值' into (cn, en)."""
    if not text:
        return "", ""
    parts = re.split(r"[\s]+", text.strip(), maxsplit=1)
    if len(parts) == 2:
        first, second = parts
        if re.search(r"[\u4e00-\u9fff]", first):
            return first, second
        return second, first
    return text.strip(), ""


def is_number(val: Any) -> bool:
    if val is None:
        return False
    try:
        float(val)
        return True
    except (ValueError, TypeError):
        return False


def _detect_table_orientation(rows: list[list[Any]]) -> str:
    """Detect if table is cross-section (region×indicator) or time-series (year×indicator)."""
    if not rows:
        return "unknown"
    first_col_values = [str(r[0]).strip() if r and r[0] else "" for r in rows[1:4]]
    year_count = sum(1 for v in first_col_values if re.match(r"^\d{4}$", v))
    region_count = sum(1 for v in first_col_values if any(k in v for k in _REGION_KEYWORDS))
    if year_count > region_count:
        return "time_series"
    return "cross_section"


def _find_header_row(rows: list[list[Any]]) -> int:
    """在前 8 行中选"最像表头"的行(带单位/首列为地区/年份关键词)。

    年鉴 xlsx 前几行常为标题(中英文)或说明, 直接取 rows[0] 会把表头当数据丢失整表。
    """
    best_idx, best_score = 0, -1.0
    for i in range(min(8, len(rows))):
        cells = [str(c).strip() if c is not None else "" for c in rows[i]]
        score = 0.0
        if any(extract_unit(c) for c in cells):
            score += 3.0
        if re.match(r"^(地区|年份|时间|园区|单位|指标|项目|Year|Region|Time)$", cells[0]):
            score += 2.5
        elif re.search(r"[\u4e00-\u9fff]", cells[0]) and not is_number(cells[0]):
            score += 1.0
        if is_number(cells[0]):
            score -= 2.0
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def _try_merge_subheader(row_after: list[Any]) -> bool:
    """rows[h+1] 是否为多级表头子行: 首列空且其余列无数字。"""
    if not row_after:
        return False
    cells = [str(c).strip() if c is not None else "" for c in row_after]
    if cells[0]:
        return False
    if any(is_number(c) and c for c in cells[1:]):
        return False
    return any(c for c in cells[1:])


def parse_xlsx(file_path: str | Path, source_label: str = "") -> list[Record]:
    """Parse an xlsx file into seven-tuple records.

    Args:
        file_path: Path to the .xlsx file.
        source_label: Source identifier for provenance.

    Returns:
        List of Record seven-tuples.
    """
    if not HAS_OPENPYXL:
        return []

    file_path = Path(file_path)
    if not source_label:
        source_label = file_path.stem

    wb = load_workbook(file_path, data_only=True, read_only=True)
    records: list[Record] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue

        # 表头搜索(跳过标题行) + 多级表头子行合并
        header_idx = _find_header_row(rows)
        header_row = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
        data_start = header_idx + 1
        if data_start < len(rows) and _try_merge_subheader(rows[data_start]):
            sub = [str(c).strip() if c is not None else "" for c in rows[data_start]]
            for i, (h, s) in enumerate(zip(header_row, sub)):
                if not h and s:
                    header_row[i] = s
            data_start += 1
        data_rows = rows[data_start:]

        orientation = _detect_table_orientation(rows[header_idx:] + data_rows)
        unit = ""

        for cell in header_row:
            u = extract_unit(cell)
            if u:
                unit = u
                break

        for row in data_rows:
            if not row or not row[0]:
                continue
            first_val = str(row[0]).strip()
            if not first_val or first_val in ("nan", "None"):
                continue

            if orientation == "time_series":
                year_match = re.match(r"^\d{4}$", first_val)
                if not year_match:
                    continue
                time_val = first_val
                space_val = "全国"
            else:
                # 截面表: 首列含中文(省/市/园区名)即认定为空间, 不再限定省级关键词
                if not re.search(r"[\u4e00-\u9fff]", first_val) or len(first_val) < 2:
                    continue
                time_val = ""
                space_val = first_val

            for col_idx, cell in enumerate(row[1:], start=1):
                if col_idx >= len(header_row):
                    break
                header = header_row[col_idx]
                if not header or header in ("nan", "None"):
                    continue

                cn_name, _ = split_cn_en(header)
                indicator = re.sub(r"[①②③④⑤⑥⑦⑧⑨⑩*★]", "", cn_name or header).strip()
                col_unit = extract_unit(header) or unit

                if is_number(cell):
                    val = float(cell)
                    src = f"xlsx:{source_label}/{sheet_name}"
                    rec = Record(
                        time=time_val,
                        space=space_val,
                        value=val,
                        unit=col_unit,
                        indicator=indicator,
                        source=src,
                        note=f"pipeline=xlsx_robust",
                    )
                    records.append(rec)

    wb.close()
    return records
