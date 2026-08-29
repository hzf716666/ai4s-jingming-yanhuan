"""Pipeline for Word (.docx) and PowerPoint (.pptx) parsing.

Uses python-docx and python-pptx — lightweight, pure-Python, no system deps.
Extracts tables from Word documents and slides from PowerPoint.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from src.schema import Record

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


def _is_number(val: Any) -> bool:
    if val is None:
        return False
    s = str(val).strip().replace(",", "").replace("，", "")
    try:
        float(s)
        return True
    except (ValueError, TypeError):
        return False


def _to_number(val: Any) -> float:
    s = str(val).strip().replace(",", "").replace("，", "")
    return float(s)


def _extract_year(text: str) -> str:
    m = re.search(r"(20\d{2}|19\d{2})", text)
    return m.group(1) if m else ""


def parse_docx(file_path: str | Path, source_label: str = "") -> list[Record]:
    """Parse Word document — extract tables as seven-tuple records.

    Args:
        file_path: Path to .docx file.
        source_label: Source identifier.

    Returns:
        List of Record seven-tuples.
    """
    if not HAS_DOCX:
        return []

    file_path = Path(file_path)
    if not source_label:
        source_label = file_path.stem

    records: list[Record] = []
    try:
        doc = DocxDocument(str(file_path))

        for table_idx, table in enumerate(doc.tables):
            if len(table.rows) < 2:
                continue
            header = [cell.text.strip() for cell in table.rows[0].cells]

            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                if not cells or not cells[0]:
                    continue
                first_val = cells[0]

                time_val = _extract_year(first_val)
                space_val = first_val if not time_val else "全国"

                for col_idx in range(1, len(cells)):
                    if col_idx >= len(header):
                        break
                    indicator = header[col_idx]
                    cell = cells[col_idx]
                    if not _is_number(cell):
                        continue
                    val = _to_number(cell)
                    records.append(Record(
                        time=time_val,
                        space=space_val,
                        value=val,
                        unit="",
                        indicator=indicator,
                        source=f"docx:{source_label}/t{table_idx}",
                        note="pipeline=docx",
                    ))
    except Exception:
        pass
    return records


def parse_pptx(file_path: str | Path, source_label: str = "") -> list[Record]:
    """Parse PowerPoint — extract tables from slides as seven-tuple records.

    Args:
        file_path: Path to .pptx file.
        source_label: Source identifier.

    Returns:
        List of Record seven-tuples.
    """
    if not HAS_PPTX:
        return []

    file_path = Path(file_path)
    if not source_label:
        source_label = file_path.stem

    records: list[Record] = []
    try:
        prs = Presentation(str(file_path))

        for slide_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if not shape.has_table:
                    continue
                table = shape.table
                if len(table.rows) < 2:
                    continue

                header = [cell.text.strip() for cell in table.rows[0].cells]

                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    if not cells or not cells[0]:
                        continue
                    first_val = cells[0]

                    time_val = _extract_year(first_val)
                    space_val = first_val if not time_val else "全国"

                    for col_idx in range(1, len(cells)):
                        if col_idx >= len(header):
                            break
                        indicator = header[col_idx]
                        cell = cells[col_idx]
                        if not _is_number(cell):
                            continue
                        val = _to_number(cell)
                        records.append(Record(
                            time=time_val,
                            space=space_val,
                            value=val,
                            unit="",
                            indicator=indicator,
                            source=f"pptx:{source_label}/s{slide_idx}",
                            note="pipeline=pptx",
                        ))
    except Exception:
        pass
    return records
